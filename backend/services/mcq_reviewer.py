"""MCQ Reviewer — the Utils tab's first tool.

A setter uploads a bulk-MCQ workbook in the exact HackerEarth template
format. We validate it (columns present, in the exact order), then review it
in stages, each one doing only what it is reliable at:

0. Deterministic checks in Python — blank/duplicate options, correct-option
   number out of range, tags missing from the taxonomy, correct-option length
   giveaway, "refer to the image" stems, exact duplicate questions. Identical
   on every run, free, and never skimmed.
1. Blind solve — two independent solvers (OpenAI and Anthropic) work each
   question WITHOUT seeing the answer key and judge every option. Python then
   compares their result to the key. This is what catches wrong keys, two
   correct options, no correct option, code bugs and contradictory puzzles.
   The model is never asked "is this key right?" — shown the key, it anchors
   on it and confirms.
2. Quality review — one entry for EVERY row (ambiguity, answer given away,
   distractor quality, bias, difficulty, ...). A row with nothing to report is
   "gtg", so a missing entry is detectable instead of looking like a pass.
3. Duplicates — exact duplicates (stem + option set, order-insensitive) in
   Python; semantic candidates (token overlap + embeddings over stem AND
   options) are adjudicated by the LLM.

Findings are tiered. Critical (wrong key, duplicate, structural, ...) fail the
row; minor (spelling, tags, formatting, ...) pass it with suggestions. The
set-level summary is computed here from the parsed rows and findings, not
trusted from the model.
"""

from __future__ import annotations

import asyncio
import io
import json
import logging
import math
import os
import re
import statistics
from dataclasses import dataclass, field

import openpyxl
import xlrd
from anthropic import AsyncAnthropic
from langsmith import traceable
from langsmith.wrappers import wrap_openai
from openai import AsyncOpenAI
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel

from core.config import settings
from core.database import Session
from core.orm import AuditLog, McqReviewJob
from schemas.mcq_review import (
    CheckResult,
    MCQReviewResult,
    OverallResults,
    QuestionReview,
    SetSummary,
)
from services import taxonomy as taxonomy_service
from services.export import safe
from utils.prompts import (
    MCQ_DUPLICATE_PROMPT,
    MCQ_REVIEWER_PROMPT,
    MCQ_SOLVER_PROMPT,
)

log = logging.getLogger(__name__)

if settings.LANGSMITH_API_KEY:
    os.environ.setdefault("LANGSMITH_API_KEY", settings.LANGSMITH_API_KEY)
    os.environ.setdefault("LANGSMITH_PROJECT", settings.LANGSMITH_PROJECT)
    os.environ.setdefault("LANGSMITH_TRACING", "true")

# The exact HackerEarth bulk-upload template header row, in this exact order.
# Reordering, inserting, or dropping a column is a hard rejection — the
# reviewer maps values positionally, not by matching header names, so a
# silently-reordered file would otherwise score the wrong field as the wrong
# check without anyone noticing.
EXPECTED_COLUMNS = (
    "Problem statement",
    "Option 1",
    "Option 2",
    "Option 3 (optional)",
    "Option 4 (optional)",
    "Correct option number (1, 2, 3, 4...etc.)",
    "Correct score",
    "Negative marking (optional)",
    "Difficulty (Easy, Medium, Hard)",
    "Timed restriction per question in seconds (optional)",
    "Skill/topic tags",
    "Partial Scoring Enabled (optional)",
    "Shuffle Options (optional)",
)

MAX_ROWS = 1000
MAX_FILE_BYTES = 10 * 1024 * 1024

# Rows per quality-review call. Every row gets an entry, so keep batches small
# enough that the model does the checks instead of skimming.
CHUNK_SIZE = 20

# Rows per blind-solve call. Smaller than CHUNK_SIZE: each row needs
# step-by-step working plus a verdict on every option.
SOLVE_CHUNK_SIZE = 10

# Batches run concurrently, capped by this semaphore rather than all at once —
# a 1000-row upload is dozens of batches, and firing them all at once trips
# OpenAI/Anthropic per-minute rate limits.
MAX_CONCURRENT_BATCHES = 5

# Two problem statements are duplicate candidates when their normalized,
# stopword-stripped token sets overlap this much (Jaccard), or their
# embeddings are this close. Deliberately generous — this only narrows the set
# for LLM adjudication, so a false positive costs a few tokens while a false
# negative means a real duplicate is never checked.
DUPLICATE_CANDIDATE_THRESHOLD = 0.4
EMBEDDING_DUPLICATE_THRESHOLD = 0.85

# Max rows per duplicate-adjudication call. Clusters are packed in without
# splitting one across batches, so an oversized cluster still goes in a single
# larger call rather than being broken apart.
DUPLICATE_BATCH_SIZE = 15

# The correct option is flagged as a length giveaway when it is the longest
# option, at least this long, and this many times the median of the others.
LENGTH_GIVEAWAY_RATIO = 1.8
LENGTH_GIVEAWAY_MIN_CHARS = 25

# "not", "true", "false" are deliberately NOT stopwords: they flip the
# meaning of code and logic questions.
_STOPWORDS = frozenset(
    [
        "a",
        "an",
        "the",
        "is",
        "are",
        "was",
        "were",
        "of",
        "to",
        "in",
        "on",
        "for",
        "with",
        "which",
        "following",
        "statements",
        "statement",
        "correct",
        "about",
        "does",
        "do",
        "you",
        "your",
        "using",
        "use",
        "it",
        "this",
        "that",
        "these",
        "those",
        "be",
        "can",
        "will",
        "would",
    ]
)

# Checks that are improvement suggestions even at status=fail. Every other
# check is critical at status=fail; borderline / na / not_evaluated are always
# minor. Critical findings fail the row, minor ones pass it with suggestions.
MINOR_ONLY_CHECKS = frozenset(
    {
        "grammar_spelling",
        "industry_relevance",
        "distractor_parity",
        "option_formatting",
        "complexity",
        "skill_tags",
    }
)

# The check names the quality-review LLM may report. Anything else it sends
# (e.g. a key verdict, which it was told not to give) is dropped.
LLM_CHECKS = frozenset(
    {
        "ambiguity",
        "grammar_spelling",
        "industry_relevance",
        "non_triviality",
        "bias_fairness",
        "distractor_quality",
        "distractor_parity",
        "option_formatting",
        "complexity",
        "skill_tags",
    }
)

# Stems that point at a figure the text-only template cannot carry.
_IMAGE_REF = re.compile(
    r"\b(?:(?:refer to|see|shown in|given in|according to|based on|consider)\s+"
    r"(?:the\s+)?(?:following\s+|above\s+|below\s+)?"
    r"(?:figure|diagram|image|picture|screenshot|illustration)"
    r"|(?:the\s+)?(?:following|above|below)\s+"
    r"(?:figure|diagram|image|picture|screenshot|illustration)"
    r"|(?:figure|diagram|image|picture|screenshot|illustration)\s+"
    r"(?:below|above|shown|attached))\b",
    re.IGNORECASE,
)


class TemplateValidationError(Exception):
    """The uploaded file doesn't match the template — never worth an LLM
    call. Carries the exact reason so the frontend can show it verbatim."""


class ReviewFailedError(Exception):
    """Both OpenAI and Anthropic failed, or neither produced a schema-valid
    response. Carries the last error seen for the audit trail."""


@dataclass
class ParsedRow:
    row_number: int  # 1-based spreadsheet row, header excluded (row 2 = first question)
    problem_statement: str
    # Positional: option i is options[i - 1]. A blank cell between filled ones
    # stays "" (flagged), so the correct-option number never points at the
    # wrong option. Trailing blanks (unused optional columns) are trimmed.
    options: list[str]
    correct_option: int | None
    difficulty: str
    skill_tags: list[str] = field(default_factory=list)

    @property
    def is_structurally_valid(self) -> bool:
        return (
            bool(self.problem_statement.strip())
            and len(self.options) >= 2
            and all(o.strip() for o in self.options)
            and self.correct_option is not None
            and 1 <= self.correct_option <= len(self.options)
        )


@dataclass
class Finding:
    """One reported problem on one row, from any stage."""

    check: str
    status: str  # a CheckResult status
    reason: str
    suggestion: str = ""
    extra: dict = field(default_factory=dict)  # remaining CheckResult fields

    @property
    def critical(self) -> bool:
        return self.status == "fail" and self.check not in MINOR_ONLY_CHECKS


def _clean(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # .xls stores every number as float: "-5", not "-5.0"
    return str(v).strip() if v is not None else ""


def _to_int(v) -> int | None:
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def _split_tags(v) -> list[str]:
    return [t.strip() for t in _clean(v).split(",") if t.strip()]


def _norm(text: str) -> str:
    return " ".join(text.lower().split())


def _norm_tag(tag: str) -> str:
    return re.sub(r"[\s_-]+", " ", tag.lower()).strip()


def _token_set(text: str) -> set[str]:
    # Words/numbers plus operator runs, so `2 ** 3 ** 2` and `a != b` keep
    # their meaning instead of collapsing to digits/letters only.
    tokens = re.findall(r"[a-z0-9_]+|[*/%+\-=<>!&|^~]+", text.lower())
    return {t for t in tokens if t not in _STOPWORDS}


def _answer_text(row: ParsedRow) -> str:
    assert row.correct_option is not None
    return row.options[row.correct_option - 1]


# ── Stage 0: deterministic checks ────────────────────────────────────────────


def _structural_findings(row: ParsedRow) -> list[Finding]:
    def fail(reason: str, suggestion: str) -> Finding:
        return Finding("structural_validation", "fail", reason, suggestion)

    out: list[Finding] = []
    if not row.problem_statement.strip():
        out.append(fail("Problem statement is blank.", "Add the question text."))
    filled = [i for i, o in enumerate(row.options, 1) if o.strip()]
    if len(filled) < 2:
        out.append(
            fail(
                f"Only {len(filled)} non-empty option(s); at least 2 are required.",
                "Add at least two answer options.",
            )
        )
    for i, o in enumerate(row.options, 1):
        if not o.strip():
            out.append(
                fail(
                    f"Option {i} is blank while a later option is filled.",
                    f"Fill option {i} or move the later options up and fix the correct option number.",
                )
            )
    k = row.correct_option
    if k is None:
        out.append(
            fail(
                "Correct option number is missing or not numeric.",
                "Enter the number (1-4) of the correct option.",
            )
        )
    elif not 1 <= k <= len(row.options):
        out.append(
            fail(
                f"Correct option number {k} is not one of the {len(row.options)} options.",
                f"Set the correct option number to a value between 1 and {len(row.options)}.",
            )
        )
    elif not row.options[k - 1].strip():
        out.append(
            fail(
                f"Correct option number {k} points at a blank option.",
                f"Fill option {k} with the correct answer or fix the number.",
            )
        )

    seen: dict[str, int] = {}
    for i, o in enumerate(row.options, 1):
        n = _norm(o)
        if not n:
            continue
        if n in seen:
            out.append(
                fail(
                    f"Options {seen[n]} and {i} are identical ({o!r}).",
                    f"Replace option {i} with a distinct, plausible distractor.",
                )
            )
        else:
            seen[n] = i
    return out


def _length_hints(row: ParsedRow) -> dict:
    lens = [len(o) for o in row.options]
    hints: dict = {"option_lengths": lens}
    k = row.correct_option
    if k is not None and 1 <= k <= len(lens) and len(lens) >= 2:
        others = lens[: k - 1] + lens[k:]
        hints["correct_is_longest"] = lens[k - 1] > max(others)
        hints["correct_vs_median_other"] = round(
            lens[k - 1] / max(statistics.median(others), 1), 2
        )
    return hints


def deterministic_findings(row: ParsedRow, taxonomy: set[str] | None) -> list[Finding]:
    """Everything decidable without an LLM. Structural problems return alone —
    a structurally invalid row gets no quality judgments."""
    structural = _structural_findings(row)
    if structural:
        return structural

    out: list[Finding] = []
    if _IMAGE_REF.search(row.problem_statement):
        out.append(
            Finding(
                "image_needed",
                "fail",
                "Image needed: the stem refers to a figure/diagram that isn't in the text.",
                "Embed the figure on the platform, or describe it in text so the question is self-contained.",
            )
        )
        return out  # nothing else can be judged without the image

    if not row.skill_tags:
        out.append(
            Finding(
                "skill_tags",
                "fail",
                "Skill/topic tags are blank.",
                "Add relevant tags from the HackerEarth taxonomy.",
                {"tag_status": "tags_not_provided", "provided_tags": []},
            )
        )
    elif taxonomy:
        unknown = [t for t in row.skill_tags if _norm_tag(t) not in taxonomy]
        if unknown:
            out.append(
                Finding(
                    "skill_tags",
                    "fail",
                    f"Tags not in the HackerEarth taxonomy: {', '.join(unknown)}.",
                    "Replace them with tags from the HackerEarth taxonomy.",
                    {"tag_status": "invalid_tags", "provided_tags": row.skill_tags},
                )
            )

    hints = _length_hints(row)
    k = row.correct_option
    assert k is not None
    correct_len = len(row.options[k - 1])
    if (
        hints.get("correct_is_longest")
        and correct_len >= LENGTH_GIVEAWAY_MIN_CHARS
        and hints["correct_vs_median_other"] >= LENGTH_GIVEAWAY_RATIO
    ):
        out.append(
            Finding(
                "distractor_parity",
                "fail",
                f"Correct option {k} is {hints['correct_vs_median_other']}x the median "
                "length of the other options — its length gives it away.",
                f"Shorten option {k} or lengthen the distractors so all options are similar in length.",
            )
        )
    return out


def _blocked(findings: list[Finding]) -> bool:
    """Rows that can't be meaningfully sent to an LLM."""
    return any(f.check in ("structural_validation", "image_needed") for f in findings)


# ── Stage 1: blind solve ─────────────────────────────────────────────────────


class _OptionVerdict(BaseModel):
    option: int
    verdict: str  # correct | incorrect | unverified; anything else = unverified


class _Solution(BaseModel):
    row_number: int
    working: str = ""
    options: list[_OptionVerdict] = []
    defect_type: str = "none"
    defect: str = ""


class _Solutions(BaseModel):
    solutions: list[_Solution] = []


_DEFECT_CHECKS = {
    "code_bug": "code_bug",
    "inconsistent_conditions": "logic_consistency",
    "image_needed": "image_needed",
}


def _classify_solution(row: ParsedRow, sol: _Solution) -> tuple[str, list[int]]:
    """Compare ONE solver's per-option verdicts to the key.
    Returns (kind, options the solver found correct); kind is one of
    ok | wrong_key | multiple_correct | none_correct | unverified."""
    assert row.correct_option is not None
    given = {o.option: o.verdict for o in sol.options}
    v = {
        i: given[i] if given.get(i) in ("correct", "incorrect") else "unverified"
        for i in range(1, len(row.options) + 1)
    }
    correct = [i for i, x in v.items() if x == "correct"]
    unverified = any(x == "unverified" for x in v.values())
    key_verdict = v[row.correct_option]
    if key_verdict == "correct":
        return ("multiple_correct" if len(correct) > 1 else "ok"), correct
    if key_verdict == "incorrect":
        if correct:
            return "wrong_key", correct
        return ("unverified" if unverified else "none_correct"), correct
    return "unverified", correct


def _opt(row: ParsedRow, i: int) -> str:
    return f"option {i} ({row.options[i - 1]!r})"


def _solver_findings(
    row: ParsedRow, solutions: list[_Solution | None]
) -> list[Finding]:
    """Turn the solvers' independent results into findings. A defect is
    critical only when every solver that answered agrees on it; when they
    disagree it is surfaced as UNVERIFIED (minor) instead of silently passing
    or failing on one model's opinion."""
    assert row.correct_option is not None
    key = row.correct_option
    responses = [s for s in solutions if s is not None]
    if not responses:
        return [
            Finding(
                "factual_accuracy",
                "na",
                "UNVERIFIED: no solver returned a result for this row, so the answer key was not checked.",
                "Verify the answer key manually.",
            )
        ]

    findings: list[Finding] = []

    # Defects in the question itself (code bug, contradictory puzzle, image).
    for dtype, check in _DEFECT_CHECKS.items():
        flagged = [s for s in responses if s.defect_type == dtype]
        if not flagged:
            continue
        detail = (flagged[0].defect or dtype.replace("_", " ")).strip()
        if len(flagged) == len(responses):
            findings.append(
                Finding(
                    check,
                    "fail",
                    ("Image needed: " if dtype == "image_needed" else "") + detail,
                    "Provide the missing image/context."
                    if dtype == "image_needed"
                    else "Fix the question so it is correct and self-consistent.",
                )
            )
        else:
            findings.append(
                Finding(
                    check,
                    "borderline",
                    f"Flagged by {len(flagged)} of {len(responses)} solvers: {detail}",
                    "Confirm whether this is a real defect.",
                )
            )
    if any(f.check == "image_needed" and f.critical for f in findings):
        return findings  # the solvers could only have guessed the answer

    results = [_classify_solution(row, s) for s in responses]
    sigs = {(k, tuple(c) if k == "wrong_key" else ()) for k, c in results}
    kinds = [k for k, _ in results]
    kind, correct = results[0]
    working = (responses[0].working or "").strip()[:300]
    work = f" Solver working: {working}" if working else ""

    if len(sigs) == 1 and kind not in ("ok", "unverified"):
        if kind == "wrong_key":
            c = correct[0]
            findings.append(
                Finding(
                    "factual_accuracy",
                    "fail",
                    f"Answer key is {_opt(row, key)} but independent solvers found "
                    f"{_opt(row, c)} correct.{work}",
                    f"Change the correct option to {c} ({row.options[c - 1]}), or fix option {key} if it was meant to be right.",
                )
            )
        elif kind == "multiple_correct":
            extra = [i for i in correct if i != key]
            findings.append(
                Finding(
                    "multiple_correct",
                    "fail",
                    "More than one option is correct: "
                    + ", ".join(_opt(row, i) for i in correct)
                    + ".",
                    f"Replace {', '.join(f'option {i}' for i in extra) or 'the extra correct option(s)'} "
                    "with a clearly incorrect distractor, or make the question multi-select.",
                )
            )
        else:  # none_correct
            findings.append(
                Finding(
                    "no_correct_answer",
                    "fail",
                    f"None of the options is correct.{work}",
                    "Add the correct answer as one of the options (or fix the stem), and set the correct option number.",
                )
            )
    elif any(k not in ("ok", "unverified") for k in kinds):
        findings.append(
            Finding(
                "factual_accuracy",
                "na",
                "UNVERIFIED: independent solvers disagree about the answer key ("
                + "; ".join(
                    f"{k}{' -> option ' + ','.join(map(str, c)) if c else ''}"
                    for k, c in results
                )
                + ").",
                "Confirm the correct answer manually before using this question.",
            )
        )
    elif all(k == "unverified" for k in kinds):
        findings.append(
            Finding(
                "factual_accuracy",
                "na",
                "UNVERIFIED: the answer could not be confirmed from the information given.",
                "Confirm the answer key manually.",
            )
        )
    return findings


def _solver_input(rows: list[ParsedRow]) -> str:
    # No key, no difficulty, no tags: nothing that could anchor the solver.
    return json.dumps(
        [
            {
                "row_number": r.row_number,
                "problem_statement": r.problem_statement,
                "options": r.options,
            }
            for r in rows
        ]
    )


# ── Stage 2: quality review ──────────────────────────────────────────────────


class _LLMRowReview(BaseModel):
    row_number: int
    comment: str = ""
    checks: dict[str, CheckResult] = {}
    suggestion: str = ""


class _LLMRowReviews(BaseModel):
    question_reviews: list[_LLMRowReview] = []


def _review_input(rows: list[ParsedRow]) -> str:
    return json.dumps(
        [
            {
                "row_number": r.row_number,
                "problem_statement": r.problem_statement,
                "options": r.options,
                "correct_option": r.correct_option,
                "difficulty": r.difficulty or None,
                "skill_tags": r.skill_tags,
                "hints": _length_hints(r),
            }
            for r in rows
        ]
    )


def _llm_findings(rr: _LLMRowReview, taxonomy_available: bool) -> list[Finding]:
    out: list[Finding] = []
    for name, c in rr.checks.items():
        if name not in LLM_CHECKS or c.status not in ("fail", "borderline"):
            continue
        if name == "skill_tags" and not taxonomy_available:
            continue
        extra = c.model_dump(exclude={"status", "reason"}, exclude_none=True)
        out.append(Finding(name, c.status, c.reason, "", extra))
    if out and rr.suggestion:
        out[0].suggestion = rr.suggestion
    return out


# ── LLM plumbing ─────────────────────────────────────────────────────────────


def _strip_fences(text: str) -> str:
    t = text.strip()
    if t.startswith("```"):
        t = t.split("\n", 1)[1] if "\n" in t else t
        t = t.rsplit("```", 1)[0]
    return t.strip()


@traceable(name="mcq_reviewer.openai", run_type="llm")
async def _call_openai(system: str, user: str) -> str:
    if not settings.OPENAI_API_KEY:
        raise RuntimeError("OPENAI_API_KEY is not configured")
    client = wrap_openai(AsyncOpenAI(api_key=settings.OPENAI_API_KEY))
    resp = await client.chat.completions.create(
        model=settings.OPENAI_MODEL,
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
        response_format={"type": "json_object"},
        temperature=0,
    )
    return resp.choices[0].message.content or ""


@traceable(name="mcq_reviewer.anthropic", run_type="llm")
async def _call_anthropic(system: str, user: str) -> str:
    if not settings.ANTHROPIC_API_KEY:
        raise RuntimeError("ANTHROPIC_API_KEY is not configured")
    # Not wrapped: langsmith's wrap_anthropic touches `client.completions`,
    # which anthropic 1.x removed. @traceable above already traces this call.
    client = AsyncAnthropic(api_key=settings.ANTHROPIC_API_KEY)
    resp = await client.messages.create(
        model=settings.ANTHROPIC_MODEL,
        max_tokens=8192,
        system=system,
        messages=[{"role": "user", "content": user}],
    )
    return "".join(b.text for b in resp.content if b.type == "text")


_PROVIDERS = {"openai": _call_openai, "anthropic": _call_anthropic}


async def _ask(provider: str, system: str, user: str, model_cls):
    raw = await _PROVIDERS[provider](system, user)
    return model_cls.model_validate_json(_strip_fences(raw))


async def _ask_with_fallback(system: str, user: str, model_cls):
    """OpenAI first; Anthropic on any failure (network, rate limit, invalid
    JSON, schema mismatch). Both failing raises ReviewFailedError — never a
    silently partial or fabricated result."""
    last_error: Exception | None = None
    for provider in ("openai", "anthropic"):
        try:
            return await _ask(provider, system, user, model_cls)
        except Exception as e:
            log.warning("mcq_reviewer: %s failed: %s", provider, e)
            last_error = e
    raise ReviewFailedError(
        f"Both OpenAI and Anthropic failed to produce a valid review: {last_error}"
    )


def _filter_valid_reviews(rows: list[ParsedRow], question_reviews: list) -> list:
    """Drop any review whose row_number isn't one we actually sent. A
    hallucination or a row_number an injected instruction in some cell
    talked the model into fabricating would otherwise corrupt the
    deterministic counts in _build_summary or show a nonexistent row on the
    frontend's flagged-questions list."""
    valid_row_numbers = {r.row_number for r in rows}
    kept = [qr for qr in question_reviews if qr.row_number in valid_row_numbers]
    if len(kept) != len(question_reviews):
        log.warning(
            "mcq_reviewer: dropped %d review(s) for row_numbers not in the uploaded file",
            len(question_reviews) - len(kept),
        )
    return kept


def _chunks(rows: list[ParsedRow], size: int) -> list[list[ParsedRow]]:
    return [rows[i : i + size] for i in range(0, len(rows), size)]


async def _solve_batch(rows: list[ParsedRow]) -> list[dict[int, _Solution]]:
    """Both providers solve the same batch independently; whichever answered
    is returned. Both failing raises."""
    user = _solver_input(rows)
    valid = {r.row_number for r in rows}

    async def one(provider: str) -> dict[int, _Solution] | None:
        try:
            parsed = await _ask(provider, MCQ_SOLVER_PROMPT, user, _Solutions)
            return {s.row_number: s for s in parsed.solutions if s.row_number in valid}
        except Exception as e:
            log.warning("mcq_reviewer: %s solver failed: %s", provider, e)
            return None

    results = await asyncio.gather(one("openai"), one("anthropic"))
    answered = [r for r in results if r is not None]
    if not answered:
        raise ReviewFailedError("Both OpenAI and Anthropic failed to solve a batch.")
    return answered


async def _review_rows(
    rows: list[ParsedRow], system: str, taxonomy_available: bool
) -> dict[int, list[Finding]]:
    """One quality-review call. Rows the model silently skipped are retried
    once; still missing, they are reported rather than passed."""
    out: dict[int, list[Finding]] = {}
    pending = rows
    for _attempt in range(2):
        if not pending:
            break
        parsed = await _ask_with_fallback(
            system, _review_input(pending), _LLMRowReviews
        )
        for rr in _filter_valid_reviews(pending, parsed.question_reviews):
            out[rr.row_number] = _llm_findings(rr, taxonomy_available)
        pending = [r for r in pending if r.row_number not in out]
    for r in pending:
        out[r.row_number] = [
            Finding(
                "review_incomplete",
                "not_evaluated",
                "The reviewer did not return a result for this row.",
                "Re-run the review for this question.",
            )
        ]
    return out


# ── Stage 3: duplicates ──────────────────────────────────────────────────────


class _DupGroup(BaseModel):
    rows: list[int]
    kind: str = "duplicate"
    reason: str = ""
    suggestion: str = ""


class _DupGroups(BaseModel):
    groups: list[_DupGroup] = []


def _dup_findings(
    group_rows: list[int], kind: str, reason: str, suggestion: str
) -> dict[int, Finding]:
    """The same finding on EVERY row of the group, each naming the others."""
    prefix = (
        "Conflicting answer keys for the same question: "
        if kind == "conflicting_key"
        else "Duplicate question: "
    )
    return {
        n: Finding(
            "duplicate_question",
            "fail",
            prefix + reason,
            suggestion or "Remove or rewrite one of the duplicate questions.",
            {"duplicate_of_rows": [m for m in group_rows if m != n]},
        )
        for n in group_rows
    }


def _exact_duplicate_groups(rows: list[ParsedRow]) -> list[tuple[list[ParsedRow], str]]:
    """Same stem and same option SET (order-insensitive). Same answer text →
    duplicate; different answer text → conflicting key."""
    groups: dict[tuple, list[ParsedRow]] = {}
    for r in rows:
        key = (_norm(r.problem_statement), tuple(sorted(_norm(o) for o in r.options)))
        groups.setdefault(key, []).append(r)
    out = []
    for g in groups.values():
        if len(g) < 2:
            continue
        answers = {_norm(_answer_text(r)) for r in g}
        out.append((g, "duplicate" if len(answers) == 1 else "conflicting_key"))
    return out


def _cluster(
    rows: list[ParsedRow], vectors: list[list[float]] | None
) -> list[list[ParsedRow]]:
    """Group rows into duplicate-candidate clusters (union-find) when their
    problem statements overlap (token Jaccard) or their stem+options
    embeddings are close. Clustering first means an adjudication batch only
    compares rows already plausibly linked, and packing never splits a
    suspected pair across calls."""
    sigs = [_token_set(r.problem_statement) for r in rows]
    parent = list(range(len(rows)))

    def find(x: int) -> int:
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    for i in range(len(rows)):
        for j in range(i + 1, len(rows)):
            linked = False
            if sigs[i] and sigs[j]:
                linked = (
                    len(sigs[i] & sigs[j]) / len(sigs[i] | sigs[j])
                    >= DUPLICATE_CANDIDATE_THRESHOLD
                )
            if not linked and vectors is not None:
                linked = (
                    math.sumprod(vectors[i], vectors[j])
                    >= EMBEDDING_DUPLICATE_THRESHOLD
                )
            if linked:
                parent[find(i)] = find(j)

    clusters: dict[int, list[ParsedRow]] = {}
    for i, r in enumerate(rows):
        clusters.setdefault(find(i), []).append(r)
    return [c for c in clusters.values() if len(c) >= 2]


async def _embed(rows: list[ParsedRow]) -> list[list[float]] | None:
    """Unit-length embeddings of stem + options + answer, or None (then the
    token filter alone runs). Never fatal: this only widens the candidate set."""
    if not settings.OPENAI_API_KEY or len(rows) < 2:
        return None
    texts = [
        r.problem_statement
        + "\n"
        + "\n".join(r.options)
        + "\nAnswer: "
        + _answer_text(r)
        for r in rows
    ]
    try:
        client = wrap_openai(AsyncOpenAI(api_key=settings.OPENAI_API_KEY))
        resp = await client.embeddings.create(
            model=settings.OPENAI_EMBEDDING_MODEL, input=texts
        )
        # Normalise ourselves rather than assume the provider did.
        out = []
        for d in resp.data:
            norm = math.sqrt(math.sumprod(d.embedding, d.embedding)) or 1.0
            out.append([x / norm for x in d.embedding])
        return out
    except Exception as e:
        log.warning("mcq_reviewer: embeddings failed, using token overlap only: %s", e)
        return None


def _pack_duplicate_batches(
    clusters: list[list[ParsedRow]],
) -> list[list[ParsedRow]]:
    """Pack clusters into DUPLICATE_BATCH_SIZE-ish batches, never splitting a
    single cluster across two batches."""
    batches: list[list[ParsedRow]] = []
    current: list[ParsedRow] = []
    for cluster in clusters:
        if current and len(current) + len(cluster) > DUPLICATE_BATCH_SIZE:
            batches.append(current)
            current = []
        current.extend(cluster)
    if current:
        batches.append(current)
    return batches


def _dup_input(rows: list[ParsedRow]) -> str:
    return json.dumps(
        [
            {
                "row_number": r.row_number,
                "problem_statement": r.problem_statement,
                "options": r.options,
                "correct_option": r.correct_option,
            }
            for r in rows
        ]
    )


# ── Summary ──────────────────────────────────────────────────────────────────


def _build_summary(
    rows: list[ParsedRow],
    question_reviews: list[QuestionReview],
    taxonomy_available: bool = True,
) -> SetSummary:
    by_row = {qr.row_number: qr for qr in question_reviews}
    total = len(rows)
    failed = sum(1 for qr in question_reviews if qr.verdict == "fail")
    suggestions_only = len(question_reviews) - failed
    passed = total - failed

    structural_failures = sum(
        1 for qr in by_row.values() if "structural_validation" in qr.checks
    )

    overall = OverallResults(
        total=total,
        passed=passed,
        failed=failed,
        pass_pct=round(100 * passed / total, 2) if total else 0.0,
        fail_pct=round(100 * failed / total, 2) if total else 0.0,
        structural_failures=structural_failures,
        pass_with_suggestions=suggestions_only,
        overall_status="clean" if not question_reviews else "issues_found",
    )

    most_frequent: dict[str, dict] = {}
    for qr in by_row.values():
        for check_name in qr.checks:
            most_frequent.setdefault(check_name, {"count": 0})
            most_frequent[check_name]["count"] += 1
    for info in most_frequent.values():
        info["pct_of_reviewed"] = (
            round(100 * info["count"] / total, 2) if total else 0.0
        )

    # Only structurally valid rows (deterministic, computed from parsed data —
    # not from the model) contribute to the distribution, grouped by option count.
    valid_rows = [r for r in rows if r.is_structurally_valid]
    by_option_count: dict[int, list[ParsedRow]] = {}
    for r in valid_rows:
        by_option_count.setdefault(len(r.options), []).append(r)

    distribution = {}
    for n, group in by_option_count.items():
        expected_pct = round(100 / n, 2)
        counts = {str(i): 0 for i in range(1, n + 1)}
        for r in group:
            counts[str(r.correct_option)] += 1
        observed = {
            k: round(100 * v / len(group), 2) if group else 0.0
            for k, v in counts.items()
        }
        status = (
            "within_tolerance"
            if all(abs(v - expected_pct) <= 5 for v in observed.values())
            else "exceeds_tolerance"
        )
        distribution[f"{n}_option"] = {
            "sample_size": len(group),
            "expected_distribution": {k: expected_pct for k in counts},
            "observed_distribution": observed,
            "tolerance": "5 percentage points",
            "status": status,
        }

    mismatches = sum(
        1
        for qr in by_row.values()
        if "complexity" in qr.checks and qr.checks["complexity"].status != "na"
    )
    structurally_valid_count = len(valid_rows)
    complexity_mismatch = {
        "structurally_valid_rows_assessed": structurally_valid_count,
        "rows_with_mismatch": mismatches,
        "mismatch_pct": (
            round(100 * mismatches / structurally_valid_count, 2)
            if structurally_valid_count
            else 0.0
        ),
    }

    tags_not_provided = sum(
        1
        for qr in by_row.values()
        if qr.checks.get("skill_tags")
        and qr.checks["skill_tags"].tag_status == "tags_not_provided"
    )

    return SetSummary(
        overall_results=overall,
        most_frequent_issues=most_frequent,
        answer_choice_distribution=distribution,
        complexity_mismatch=complexity_mismatch,
        skill_tag_gaps={"rows_missing_tags": tags_not_provided},
        skill_tag_analysis_status=None
        if taxonomy_available
        else "taxonomy_not_provided",
    )


def _assemble(
    rows: list[ParsedRow],
    findings: dict[int, list[Finding]],
    taxonomy_available: bool,
) -> MCQReviewResult:
    """Findings → tiered QuestionReviews + a comment for EVERY row."""
    reviews: list[QuestionReview] = []
    comments: dict[int, str] = {}
    for r in rows:
        fs = findings.get(r.row_number, [])
        if not fs:
            comments[r.row_number] = "gtg"
            continue
        checks: dict[str, CheckResult] = {}
        for f in fs:
            if f.check in checks:  # same check from two stages: keep both reasons
                checks[f.check].reason += " " + f.reason
                continue
            checks[f.check] = CheckResult(status=f.status, reason=f.reason, **f.extra)
        critical = any(f.critical for f in fs)
        suggestion = " ".join(dict.fromkeys(f.suggestion for f in fs if f.suggestion))
        reviews.append(
            QuestionReview(
                row_number=r.row_number,
                verdict="fail" if critical else "pass_with_suggestions",
                checks=checks,
                suggestion=suggestion,
            )
        )
        reasons = " | ".join(f"{name}: {c.reason}" for name, c in checks.items())
        comments[r.row_number] = (
            ("FAIL — " if critical else "Suggestions — ")
            + reasons
            + (f" Fix: {suggestion}" if suggestion else "")
        )
    return MCQReviewResult(
        question_reviews=reviews,
        row_comments=comments,
        set_summary=_build_summary(rows, reviews, taxonomy_available),
    )


# ── Orchestration ────────────────────────────────────────────────────────────


@traceable(name="mcq_reviewer.review_set", run_type="chain")
async def review_mcqs(
    rows: list[ParsedRow],
    taxonomy_block: str,
    taxonomy_tags: set[str] | None = None,
) -> MCQReviewResult:
    """Run every stage and assemble the result. See the module docstring.
    `taxonomy_tags` is the tag list used for the deterministic membership
    check; None/empty means no taxonomy is available."""
    taxonomy = {_norm_tag(t) for t in taxonomy_tags} if taxonomy_tags else None
    taxonomy_available = bool(taxonomy)

    system = MCQ_REVIEWER_PROMPT
    if taxonomy_block:
        system += (
            "\n\n---\n\n"
            "## Skill Taxonomy Reference (hackerearth-skill-taxonomy.md)\n\n"
            + taxonomy_block
        )

    findings: dict[int, list[Finding]] = {r.row_number: [] for r in rows}

    # Stage 0 — deterministic.
    for r in rows:
        findings[r.row_number].extend(deterministic_findings(r, taxonomy))
    llm_rows = [r for r in rows if not _blocked(findings[r.row_number])]

    # Exact duplicates are also deterministic; those rows skip adjudication.
    dup_rows = [r for r in llm_rows if r.is_structurally_valid]
    exact_ids: set[int] = set()
    for group, kind in _exact_duplicate_groups(dup_rows):
        nums = [r.row_number for r in group]
        exact_ids.update(nums)
        reason = (
            f"rows {', '.join(map(str, nums))} have the same stem and the same options"
            + (" and answer." if kind == "duplicate" else " but different answers.")
        )
        for n, f in _dup_findings(nums, kind, reason, "").items():
            findings[n].append(f)

    semaphore = asyncio.Semaphore(MAX_CONCURRENT_BATCHES)

    async def bounded(coro_fn, *args):
        async with semaphore:
            return await coro_fn(*args)

    async def stage_solve() -> None:
        batches = _chunks(
            [r for r in llm_rows if r.is_structurally_valid], SOLVE_CHUNK_SIZE
        )
        results = await asyncio.gather(*(bounded(_solve_batch, b) for b in batches))
        for batch, solutions in zip(batches, results):
            for r in batch:
                findings[r.row_number].extend(
                    _solver_findings(r, [s.get(r.row_number) for s in solutions])
                )

    async def stage_review() -> None:
        batches = _chunks(llm_rows, CHUNK_SIZE)
        results = await asyncio.gather(
            *(bounded(_review_rows, b, system, taxonomy_available) for b in batches)
        )
        for res in results:
            for n, fs in res.items():
                findings[n].extend(fs)

    async def stage_duplicates() -> None:
        candidates = [r for r in dup_rows if r.row_number not in exact_ids]
        if len(candidates) < 2:
            return
        vectors = await _embed(candidates)
        clusters = await asyncio.to_thread(_cluster, candidates, vectors)
        batches = _pack_duplicate_batches(clusters)

        async def adjudicate(batch: list[ParsedRow]) -> None:
            parsed = await _ask_with_fallback(
                MCQ_DUPLICATE_PROMPT, _dup_input(batch), _DupGroups
            )
            valid = {r.row_number for r in batch}
            claimed: set[int] = set()
            for g in parsed.groups:
                nums = list(dict.fromkeys(g.rows))
                if len(nums) < 2 or not set(nums) <= valid or claimed & set(nums):
                    continue  # hallucinated row, or a row in two groups
                claimed.update(nums)
                for n, f in _dup_findings(nums, g.kind, g.reason, g.suggestion).items():
                    findings[n].append(f)

        await asyncio.gather(*(bounded(adjudicate, b) for b in batches))

    await asyncio.gather(stage_solve(), stage_review(), stage_duplicates())
    return _assemble(rows, findings, taxonomy_available)


# ── Workbook I/O ─────────────────────────────────────────────────────────────


def _rows_from_sheet(header: list[str], data_rows: list[tuple]) -> list[ParsedRow]:
    validate_columns(header)
    parsed = []
    for i, r in enumerate(data_rows):
        row_number = i + 2  # header is row 1
        options = [_clean(r[j]) for j in (1, 2, 3, 4)]
        while options and not options[-1]:
            options.pop()  # unused optional columns, not blank distractors
        parsed.append(
            ParsedRow(
                row_number=row_number,
                problem_statement=_clean(r[0]),
                options=options,
                correct_option=_to_int(r[5]),
                difficulty=_clean(r[8]),
                skill_tags=_split_tags(r[10]),
            )
        )
    return parsed


def validate_columns(header: list[str]) -> None:
    """Position-by-position match against EXPECTED_COLUMNS. Reordering,
    inserting, or removing a column fails here — before any LLM call."""
    got = [_clean(h) for h in header]
    if len(got) < len(EXPECTED_COLUMNS):
        raise TemplateValidationError(
            f"Expected {len(EXPECTED_COLUMNS)} columns, found {len(got)}."
        )
    for i, expected in enumerate(EXPECTED_COLUMNS):
        actual = got[i] if i < len(got) else "(missing)"
        if actual != expected:
            raise TemplateValidationError(
                f"Column {i + 1} must be {expected!r} in this exact position, "
                f"found {actual!r}. Columns cannot be reordered, renamed, "
                "inserted, or removed."
            )


def _read_raw_rows(content: bytes, filename: str) -> tuple[list, list[tuple]]:
    """Header row + data rows, straight off the sheet with no cleaning —
    shared by parse_workbook (which derives ParsedRow from it) and
    build_reviewed_workbook (which needs every original column verbatim to
    append Pass/Fail/Suggestion onto)."""
    if len(content) > MAX_FILE_BYTES:
        raise TemplateValidationError(
            f"File is larger than the {MAX_FILE_BYTES // (1024 * 1024)}MB limit."
        )

    lower = filename.lower()
    try:
        if lower.endswith(".xlsx"):
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            ws = wb.worksheets[0]
            rows = list(ws.iter_rows(values_only=True))
        elif lower.endswith(".xls"):
            wb = xlrd.open_workbook(file_contents=content)
            ws = wb.sheet_by_index(0)
            rows = [
                tuple(ws.cell_value(r, c) for c in range(ws.ncols))
                for r in range(ws.nrows)
            ]
        else:
            raise TemplateValidationError(
                "Unsupported file type — upload the .xls or .xlsx bulk template."
            )
    except TemplateValidationError:
        raise
    except Exception as e:
        raise TemplateValidationError(f"Could not read the file: {e}")

    if not rows:
        raise TemplateValidationError("The file is empty.")
    header, data_rows = list(rows[0]), rows[1:]
    if not data_rows:
        raise TemplateValidationError("The file has a header row but no questions.")
    if len(data_rows) > MAX_ROWS:
        raise TemplateValidationError(
            f"{len(data_rows)} questions found — the limit is {MAX_ROWS} per upload."
        )
    return header, data_rows


def parse_workbook(content: bytes, filename: str) -> list[ParsedRow]:
    header, data_rows = _read_raw_rows(content, filename)
    return _rows_from_sheet(header, data_rows)


RED_FILL = PatternFill("solid", fgColor="FFC7CE")
RED_FONT = Font(color="9C0006")
BOLD_HEADER = Font(bold=True)


def build_reviewed_workbook(
    content: bytes, filename: str, result: MCQReviewResult
) -> bytes:
    """The original uploaded sheet, unchanged column-for-column, with Pass,
    Fail, and Suggestion appended. Every row gets a Suggestion: "gtg" or the
    findings. Only critical findings (verdict "fail") mark Fail=Yes (red);
    suggestions-only rows pass.

    Original cells are written verbatim. Formula injection is neutralised by
    storing strings as text (not by prefixing an apostrophe, which would
    corrupt values like "-5"); only our own generated text goes through
    safe()."""
    header, data_rows = _read_raw_rows(content, filename)
    reviews_by_row = {qr.row_number: qr for qr in result.question_reviews}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reviewed"

    new_header = [*header, "Pass", "Fail", "Suggestion"]
    for col, label in enumerate(new_header, 1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = BOLD_HEADER

    pass_col = len(header) + 1
    fail_col = len(header) + 2
    suggestion_col = len(header) + 3
    for i, raw_row in enumerate(data_rows):
        row_number = i + 2
        excel_row = i + 2
        for col, value in enumerate(raw_row, 1):
            cell = ws.cell(row=excel_row, column=col, value=value)
            if isinstance(value, str):
                cell.data_type = "s"  # never a formula / error literal

        qr = reviews_by_row.get(row_number)
        failed = qr is not None and qr.verdict == "fail"
        ws.cell(row=excel_row, column=pass_col, value="No" if failed else "Yes")
        fail_cell = ws.cell(
            row=excel_row, column=fail_col, value="Yes" if failed else "No"
        )
        if failed:
            fail_cell.fill = RED_FILL
            fail_cell.font = RED_FONT
        comment = result.row_comments.get(row_number) or (
            qr.suggestion if qr else "gtg"
        )
        ws.cell(row=excel_row, column=suggestion_col, value=safe(comment))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def run_review_job(job_id: int, content: bytes) -> None:
    """The background task kicked off by POST /api/utils/mcq-reviewer. Opens
    its own DB session — the request's session is gone by the time a
    BackgroundTasks callback runs (same pattern as services.content_issues.sync).
    Every stage transition is persisted on the job row and logged to AuditLog,
    so a failure anywhere still leaves a clear, queryable trail of what was
    uploaded, by whom, and whether it finished."""
    async with Session() as db:
        job = await db.get(McqReviewJob, job_id)
        if job is None:
            return

        async def _stage(status: str, **payload) -> None:
            job.status = status
            await db.commit()
            db.add(
                AuditLog(
                    user_id=job.user_id,
                    action=f"mcq_reviewer.{status}",
                    entity_type="mcq_review_job",
                    entity_id=str(job_id),
                    payload={"filename": job.filename, **payload},
                )
            )
            await db.commit()

        try:
            await _stage("parsing")
            rows = parse_workbook(content, job.filename)
            await _stage("parsed", row_count=len(rows))

            await _stage("reviewing")
            taxonomy_block = await taxonomy_service.taxonomy_prompt_block(db)
            taxonomy_tags = {t.tag for t in await taxonomy_service.list_taxonomy(db)}
            result = await review_mcqs(
                rows,
                taxonomy_block,
                taxonomy_tags,
                langsmith_extra={
                    "name": f"mcq_reviewer.review_set[{job.filename}]",
                    "metadata": {
                        "job_id": job_id,
                        "filename": job.filename,
                        "row_count": len(rows),
                    },
                },
            )

            job.result = result.model_dump(mode="json")
            await _stage("done", row_count=len(rows))
        except TemplateValidationError as e:
            job.error = str(e)
            await _stage("failed", error=str(e), reason="template_validation")
        except ReviewFailedError as e:
            job.error = str(e)
            await _stage("failed", error=str(e), reason="llm_failure")
        except Exception as e:  # pragma: no cover - unexpected, never leave a job stuck
            log.exception("mcq_reviewer: job %s crashed", job_id)
            job.error = f"Unexpected error: {e}"
            await _stage("failed", error=str(e), reason="unexpected")
