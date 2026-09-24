"""Shared parsing/classification logic for Event Question Review data
imports — used by both the one-off local seed script
(scripts/seed_event_review.py, which truncates and rebuilds) and the
admin-facing upload flow (api/utils_routes.py's /event-review/import
routes, which is additive and skips conflicts). See EVENT_QUESTION_REVIEW.md.

`classify_l1`/`classify_l2` are pure (no DB, no mutation) so the same logic
drives both the preview (what WOULD happen) and the actual write
(`apply_l1`/`apply_l2`, which mutate a `QuestionReview`) — a preview that
could drift from what a confirm actually does would be worse than no
preview at all.
"""

from __future__ import annotations

import html
import io
import logging
from dataclasses import dataclass
from datetime import UTC, datetime

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import undefer

from core.database import Session
from core.orm import EventReviewImportJob, Member, QuestionReview, QuestionReviewEvent

log = logging.getLogger(__name__)

L1_VERDICT_MAP = {"GTG": "no_issue_found", "Fixed": "fixed"}
# "Partial pass" has no equivalent in the current 3-verdict model (product
# decision: fold into no_issue_found). Plain "GTG" at L2 is just confirmation
# of whatever L1 already decided, so it doesn't touch last_verdict.
L2_VERDICT_MAP = {"Removed": "removed", "Partial pass": "no_issue_found"}

REQUIRED_COLUMNS = ("Problem ID", "Title", "Level", "Tags", "Score", "Description")


def clean(v) -> str:
    return html.unescape(str(v).strip()) if v not in (None, "") else ""


def discover_data_sheets(wb) -> list[str]:
    """Any sheet name can be an event slug — what makes it a *data* sheet is
    having the columns this needs, not a hardcoded name list."""
    names = []
    for name in wb.sheetnames:
        header = next(wb[name].iter_rows(min_row=1, max_row=1, values_only=True), None)
        cols = {h for h in (header or ()) if h}
        if "QuestionType" in cols and (
            "Setter Template ID" in cols or "Setter ID" in cols
        ):
            names.append(name)
    return names


def missing_columns(ws: Worksheet) -> list[str]:
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    cols = {h for h in (header or ()) if h}
    return [c for c in REQUIRED_COLUMNS if c not in cols]


def read_sheet(ws: Worksheet) -> list[dict]:
    """One dict per unique Setter Template ID. A question pooled into more
    than one section reappears with identical content but only one copy
    carries the real review columns (the other is blank) — the row with a
    status wins over first-seen."""
    header = [
        h for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True)) if h
    ]
    rows: dict[int, tuple[dict, bool]] = {}
    for raw in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(header, raw))
        stid = d.get("Setter Template ID") or d.get("Setter ID")
        if not stid:
            continue
        stid = int(stid)
        d["_stid"] = stid
        has_status = bool(
            clean(d.get("L1 review status")) or clean(d.get("L2 review status"))
        )
        if stid not in rows or (has_status and not rows[stid][1]):
            rows[stid] = (d, has_status)
    return [d for d, _ in rows.values()]


async def resolve_reviewer(
    db: AsyncSession, cache: dict[str, int | None], name: str
) -> tuple[int | None, str | None]:
    """First names only in the sheet ("Vishal") -> a real Member row. Exact
    case-insensitive match first, then a unique "starts with" match. No
    unique match -> (None, a warning) — never a fake account."""
    name = name.strip()
    if not name:
        return None, None
    key = name.lower()
    if key in cache:
        mid = cache[key]
        return mid, (None if mid else f"no member found for reviewer name {name!r}")

    exact = await db.scalar(
        select(Member).where(func.lower(Member.display_name) == key)
    )
    if exact:
        cache[key] = exact.id
        return exact.id, None

    candidates = (
        await db.scalars(select(Member).where(Member.display_name.ilike(f"{name}%")))
    ).all()
    if len(candidates) == 1:
        cache[key] = candidates[0].id
        return candidates[0].id, None

    cache[key] = None
    if len(candidates) > 1:
        return (
            None,
            f"ambiguous reviewer name {name!r} matches {[c.display_name for c in candidates]}",
        )
    return None, f"no member found for reviewer name {name!r}"


@dataclass
class L1Outcome:
    verdict: (
        str | None
    )  # None => reviewer noted but status unmapped/absent, nothing to write
    note: str
    warning: str | None = None


@dataclass
class L2Outcome:
    verdict: str | None  # None => plain confirmation, doesn't change last_verdict
    comment: str | None
    warning: str | None = None


def classify_l1(status: str, comment: str) -> L1Outcome | None:
    if not status:
        return None
    verdict = L1_VERDICT_MAP.get(status)
    if verdict is None:
        return L1Outcome(
            verdict=None, note="", warning=f"unmapped L1 review status {status!r}"
        )
    note = comment or (
        "GTG."
        if verdict == "no_issue_found"
        else f"{status} (no summary in source sheet)."
    )
    return L1Outcome(verdict=verdict, note=note)


def classify_l2(status: str, comment: str) -> L2Outcome | None:
    if not status:
        return None
    return L2Outcome(verdict=L2_VERDICT_MAP.get(status), comment=comment or None)


def apply_l1(
    qr: QuestionReview,
    reviewer_id: int | None,
    outcome: L1Outcome | None,
    when: datetime,
):
    if reviewer_id:
        qr.l1_assignee_id = reviewer_id
    if outcome is None:
        if reviewer_id:
            qr.l1_status = "in_progress"
        return None
    if outcome.verdict is None:
        return None
    qr.l1_status = "done"
    qr.l1_comments = outcome.note
    qr.last_verdict = outcome.verdict
    qr.last_reviewed_at = when
    qr.issue_summary = outcome.note if outcome.verdict == "fixed" else None
    return ("l1", reviewer_id, outcome.verdict, outcome.note)


def apply_l2(
    qr: QuestionReview,
    reviewer_id: int | None,
    outcome: L2Outcome | None,
    when: datetime,
):
    if reviewer_id:
        qr.l2_assignee_id = reviewer_id
    if outcome is None:
        if reviewer_id:
            qr.l2_status = "pending"
        return None
    qr.l2_status = "done"
    qr.l2_comments = outcome.comment
    if outcome.verdict:
        qr.last_verdict = outcome.verdict
        qr.last_reviewed_at = when
        if outcome.verdict == "removed":
            qr.removal_reason = (
                outcome.comment or "Removed (no reason recorded in source sheet)."
            )
    return ("l2", reviewer_id, outcome.verdict, outcome.comment)


def _now() -> datetime:
    return datetime.now(UTC).replace(tzinfo=None)


async def process_workbook(db: AsyncSession, content: bytes, *, dry_run: bool) -> dict:
    """One walk of the workbook drives both the preview (`dry_run=True`,
    nothing written) and the real import (`dry_run=False`) — identical
    classification/conflict logic either way, so the preview can't drift
    from what confirming it actually does.

    A row is skipped ENTIRELY (content included) if either level already has
    real review work (`l1_status`/`l2_status` == "done") and the file has
    data for that level — reported as a conflict, never silently applied.
    Re-importing the same file is idempotent: an event identical to one
    already recorded for that (question, event, level, verdict, note) is not
    re-inserted.
    """
    wb = load_workbook(io.BytesIO(content), data_only=True)
    sheet_names = discover_data_sheets(wb)
    if not sheet_names:
        raise ValueError(
            "No review sheets found — a data sheet needs a header row with a "
            "QuestionType column and a Setter Template ID (or Setter ID) column."
        )

    reviewer_cache: dict[str, int | None] = {}
    sheets: list[dict] = []
    warnings: list[str] = []
    conflicts: list[dict] = []
    now = _now()

    for slug in sheet_names:
        ws = wb[slug]
        missing = missing_columns(ws)
        rows = read_sheet(ws)

        existing_event_keys: set[tuple] = set()
        if not dry_run:
            existing = await db.execute(
                select(
                    QuestionReviewEvent.setter_template_id,
                    QuestionReviewEvent.level,
                    QuestionReviewEvent.verdict,
                    QuestionReviewEvent.note,
                ).where(QuestionReviewEvent.event_slug == slug)
            )
            existing_event_keys = set(existing.all())

        reviews_count = conflicts_count = events_count = 0

        for d in rows:
            stid = d["_stid"]
            qr = await db.get(QuestionReview, stid)

            l1_outcome = classify_l1(
                clean(d.get("L1 review status")), clean(d.get("L1 review comments"))
            )
            l2_outcome = classify_l2(
                clean(d.get("L2 review status")), clean(d.get("L2 review comments"))
            )
            if l1_outcome and l1_outcome.warning:
                warnings.append(f"{slug} #{stid}: {l1_outcome.warning}")
            if l2_outcome and l2_outcome.warning:
                warnings.append(f"{slug} #{stid}: {l2_outcome.warning}")

            l1_reviewer_id, w1 = await resolve_reviewer(
                db, reviewer_cache, clean(d.get("L1 reviewer"))
            )
            l2_reviewer_id, w2 = await resolve_reviewer(
                db, reviewer_cache, clean(d.get("L2 reviewer"))
            )
            if w1:
                warnings.append(f"{slug} #{stid}: {w1}")
            if w2:
                warnings.append(f"{slug} #{stid}: {w2}")

            has_incoming_review = bool(
                l1_outcome or l2_outcome or l1_reviewer_id or l2_reviewer_id
            )
            conflict_levels = []
            if (
                qr is not None
                and qr.l1_status == "done"
                and (l1_outcome is not None or l1_reviewer_id is not None)
            ):
                conflict_levels.append("l1")
            if (
                qr is not None
                and qr.l2_status == "done"
                and (l2_outcome is not None or l2_reviewer_id is not None)
            ):
                conflict_levels.append("l2")

            if conflict_levels:
                conflicts_count += 1
                conflicts.append(
                    {
                        "slug": slug,
                        "setter_template_id": stid,
                        "title": clean(d.get("Title")),
                        "levels": conflict_levels,
                    }
                )
                continue  # whole row skipped — never overwrite real review work

            if not has_incoming_review:
                continue  # content-only row, nothing to review, nothing to report

            if dry_run:
                reviews_count += 1
                continue

            if qr is None:
                qr = QuestionReview(setter_template_id=stid)
                db.add(qr)
            qr.question_type = clean(d.get("QuestionType"))
            qr.problem_id = int(d.get("Problem ID") or 0)
            qr.title = clean(d.get("Title"))
            qr.level = clean(d.get("Level"))
            qr.tags = [t.strip() for t in clean(d.get("Tags")).split(",") if t.strip()]
            qr.score = int(d.get("Score") or 0)
            qr.description = clean(d.get("Description"))
            if not qr.last_reviewed_slug:
                qr.last_reviewed_slug = slug

            for event in (
                apply_l1(qr, l1_reviewer_id, l1_outcome, now),
                apply_l2(qr, l2_reviewer_id, l2_outcome, now),
            ):
                if event is None:
                    continue
                level, actor_id, verdict, note = event
                key = (stid, level, verdict, note or "")
                if key in existing_event_keys:
                    continue  # identical to an already-recorded review — idempotent re-import
                db.add(
                    QuestionReviewEvent(
                        setter_template_id=stid,
                        event_slug=slug,
                        level=level,
                        actor_member_id=actor_id,
                        verdict=verdict,
                        note=note or "",
                        created_at=now,
                    )
                )
                existing_event_keys.add(key)
                events_count += 1
            reviews_count += 1

        sheets.append(
            {
                "slug": slug,
                "questions": len(rows),
                "missing_columns": missing,
                "reviews": reviews_count,
                "conflicts": conflicts_count,
                "events": events_count,
            }
        )

    return {
        "dry_run": dry_run,
        "sheets": sheets,
        "warnings": warnings,
        "conflicts": conflicts,
        "total_questions": sum(s["questions"] for s in sheets),
        "total_reviews": sum(s["reviews"] for s in sheets),
        "total_conflicts": sum(s["conflicts"] for s in sheets),
        "total_events": sum(s["events"] for s in sheets),
    }


async def run_validation(job_id: int, content: bytes) -> None:
    """Background task, phase 1 (upload): parse + classify only — dry_run,
    nothing written beyond the job row's own preview/status."""
    async with Session() as db:
        job = await db.get(EventReviewImportJob, job_id)
        if job is None:
            return
        job.status = "validating"
        await db.commit()

        try:
            preview = await process_workbook(db, content, dry_run=True)
        except Exception as e:
            log.warning(
                "event review import validation failed for job %s: %s", job_id, e
            )
            job = await db.get(EventReviewImportJob, job_id)
            job.status = "error"
            job.error = str(e)
            await db.commit()
            return

        job = await db.get(EventReviewImportJob, job_id)
        job.preview = preview
        job.status = "ready_for_review"
        await db.commit()


async def run_import(job_id: int) -> None:
    """Background task, phase 2 (confirm): the real upsert. Only runs if the
    job is still in `ready_for_review` — an admin must have seen the preview
    first; this is never reachable straight from upload."""
    async with Session() as db:
        job = await db.scalar(
            select(EventReviewImportJob)
            .options(undefer(EventReviewImportJob.source_file))
            .where(EventReviewImportJob.id == job_id)
        )
        if job is None or job.status != "ready_for_review":
            return
        job.status = "importing"
        await db.commit()

        try:
            result = await process_workbook(db, job.source_file, dry_run=False)
            await db.commit()
        except Exception as e:
            await db.rollback()
            log.warning("event review import failed for job %s: %s", job_id, e)
            job = await db.get(EventReviewImportJob, job_id)
            job.status = "error"
            job.error = str(e)
            await db.commit()
            return

        job = await db.get(EventReviewImportJob, job_id)
        job.result = result
        job.status = "done"
        await db.commit()
