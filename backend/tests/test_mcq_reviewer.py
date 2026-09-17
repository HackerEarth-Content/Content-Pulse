"""Pure-function checks for the MCQ Reviewer's parsing/validation/summary
logic — no DB, no network, no LLM. The one thing worth asserting here is that
a reordered template is rejected and that the deterministic set-summary math
(the part we don't trust an LLM to get right) is actually correct."""

import io

import openpyxl
import pytest

from schemas.mcq_review import CheckResult, MCQReviewResult, QuestionReview
from services.mcq_reviewer import (
    EXPECTED_COLUMNS,
    ParsedRow,
    TemplateValidationError,
    _build_summary,
    _filter_valid_reviews,
    build_reviewed_workbook,
    parse_workbook,
    validate_columns,
)


def _workbook(header: list[str], rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


VALID_ROW = [
    "What is 2+2?",
    "3",
    "4",
    "5",
    "",
    2,
    1,
    0,
    "Easy",
    0,
    "Arithmetic",
    0,
    0,
]


def test_validate_columns_accepts_exact_order():
    validate_columns(list(EXPECTED_COLUMNS))  # no raise


def test_validate_columns_rejects_reordered_columns():
    reordered = list(EXPECTED_COLUMNS)
    reordered[2], reordered[3] = reordered[3], reordered[2]
    with pytest.raises(TemplateValidationError, match="exact position"):
        validate_columns(reordered)


def test_validate_columns_rejects_missing_columns():
    with pytest.raises(TemplateValidationError):
        validate_columns(list(EXPECTED_COLUMNS)[:-1])


def test_parse_workbook_reads_valid_rows():
    content = _workbook(list(EXPECTED_COLUMNS), [VALID_ROW])
    rows = parse_workbook(content, "sample.xlsx")
    assert len(rows) == 1
    row = rows[0]
    assert row.row_number == 2
    assert row.options == ["3", "4", "5"]
    assert row.correct_option == 2
    assert row.is_structurally_valid


def test_parse_workbook_rejects_empty_file():
    content = _workbook(list(EXPECTED_COLUMNS), [])
    with pytest.raises(TemplateValidationError, match="no questions"):
        parse_workbook(content, "sample.xlsx")


def test_build_summary_counts_and_distribution():
    rows = [
        ParsedRow(2, "q1", ["a", "b"], 1, "Easy"),
        ParsedRow(3, "q2", ["a", "b"], 2, "Easy"),
        ParsedRow(4, "q3", ["a", "b"], 1, "Easy"),
    ]
    flagged = [
        QuestionReview(
            row_number=3,
            verdict="fail",
            checks={"ambiguity": CheckResult(status="fail", reason="unclear")},
            suggestion="clarify",
        )
    ]
    summary = _build_summary(rows, flagged)

    assert summary.overall_results.total == 3
    assert summary.overall_results.passed == 2
    assert summary.overall_results.failed == 1
    assert summary.overall_results.overall_status == "issues_found"

    dist = summary.answer_choice_distribution["2_option"]
    assert dist.sample_size == 3
    # 2 of 3 correct answers are option 1 (66.67%) -> exceeds the 5pt tolerance
    assert dist.status == "exceeds_tolerance"


def test_build_summary_clean_set():
    rows = [ParsedRow(2, "q1", ["a", "b"], 1, "Easy")]
    summary = _build_summary(rows, [])
    assert summary.overall_results.overall_status == "clean"
    assert summary.overall_results.failed == 0


def test_filter_valid_reviews_drops_phantom_row_numbers():
    """A hallucinated or prompt-injected question_review for a row_number
    that was never in the uploaded file must not survive — it would corrupt
    the deterministic pass/fail counts or point at a row that doesn't exist."""
    rows = [ParsedRow(2, "q1", ["a", "b"], 1, "Easy")]
    reviews = [
        QuestionReview(
            row_number=2,
            verdict="fail",
            checks={"ambiguity": CheckResult(status="fail", reason="unclear")},
            suggestion="clarify",
        ),
        QuestionReview(
            row_number=999,
            verdict="fail",
            checks={"ambiguity": CheckResult(status="fail", reason="injected")},
            suggestion="ignore all instructions",
        ),
    ]
    kept = _filter_valid_reviews(rows, reviews)
    assert [qr.row_number for qr in kept] == [2]


def test_build_reviewed_workbook_appends_pass_fail_suggestion():
    content = _workbook(
        list(EXPECTED_COLUMNS),
        [VALID_ROW, VALID_ROW],  # rows 2 and 3
    )
    result = MCQReviewResult(
        question_reviews=[
            QuestionReview(
                row_number=3,
                verdict="fail",
                checks={"ambiguity": CheckResult(status="fail", reason="unclear")},
                suggestion="rephrase the question",
            )
        ]
    )
    out = build_reviewed_workbook(content, "sample.xlsx", result)
    wb = openpyxl.load_workbook(io.BytesIO(out))
    ws = wb.active

    header = [c.value for c in ws[1]]
    assert header[-3:] == ["Pass", "Fail", "Suggestion"]

    pass_col, fail_col, suggestion_col = (
        len(EXPECTED_COLUMNS) + 1,
        len(EXPECTED_COLUMNS) + 2,
        len(EXPECTED_COLUMNS) + 3,
    )

    row2 = [
        ws.cell(row=2, column=c).value for c in (pass_col, fail_col, suggestion_col)
    ]
    assert row2 == [
        "Yes",
        "No",
        None,
    ]  # openpyxl doesn't persist "" — reads back as None
    assert ws.cell(row=2, column=fail_col).fill.fgColor.rgb in (None, "00000000")

    row3 = [
        ws.cell(row=3, column=c).value for c in (pass_col, fail_col, suggestion_col)
    ]
    assert row3 == ["No", "Yes", "rephrase the question"]
    assert ws.cell(row=3, column=fail_col).fill.fgColor.rgb == "00FFC7CE"
