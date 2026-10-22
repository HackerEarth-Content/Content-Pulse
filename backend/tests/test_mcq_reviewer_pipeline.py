"""Staged-pipeline checks for the MCQ Reviewer — the failure modes the team
reported (wrong keys passing, two correct options, duplicate options, missing
correct answer, length giveaway, duplicates with reordered options, minor
issues failing rows, apostrophes in the reviewed sheet). Deterministic stages
are asserted directly; the LLM stages run against fake providers so the
orchestration, comparison-to-key and severity logic are covered with no
network."""

import io
import json

import openpyxl

from schemas.mcq_review import MCQReviewResult
from services import mcq_reviewer as svc
from services.mcq_reviewer import (
    EXPECTED_COLUMNS,
    ParsedRow,
    _assemble,
    _cluster,
    _exact_duplicate_groups,
    _OptionVerdict,
    _Solution,
    _solver_findings,
    build_reviewed_workbook,
    deterministic_findings,
    review_mcqs,
)


def _row(n, stem, options, key, tags=("Python",), difficulty="Easy"):
    return ParsedRow(n, stem, list(options), key, difficulty, list(tags))


def _checks(findings):
    return {f.check for f in findings}


# ── Stage 0: deterministic ───────────────────────────────────────────────────


def test_duplicate_options_flagged():
    fs = deterministic_findings(_row(2, "Q?", ["Stack", "stack ", "Queue"], 3), None)
    assert _checks(fs) == {"structural_validation"}
    assert fs[0].critical and "identical" in fs[0].reason


def test_correct_option_out_of_range_and_missing():
    assert deterministic_findings(_row(2, "Q?", ["a", "b"], 5), None)[0].critical
    assert deterministic_findings(_row(2, "Q?", ["a", "b"], None), None)[0].critical


def test_blank_middle_option_does_not_shift_indices():
    sheet_row = ["Q?", "a", "", "c", "", 3, 1, 0, "Easy", 0, "Python", 0, 0]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(EXPECTED_COLUMNS))
    ws.append(sheet_row)
    buf = io.BytesIO()
    wb.save(buf)
    row = svc.parse_workbook(buf.getvalue(), "x.xlsx")[0]
    assert row.options == ["a", "", "c"]  # option 3 is still "c"
    assert any(
        "Option 2 is blank" in f.reason for f in deterministic_findings(row, None)
    )


def test_tags_checked_against_taxonomy_consistently():
    taxonomy = {"python", "operating systems"}
    ok = deterministic_findings(
        _row(2, "Q?", ["a", "b"], 1, tags=["Operating-Systems"]), taxonomy
    )
    assert not ok  # hyphen/case differences are formatting, not substance
    bad = deterministic_findings(
        _row(2, "Q?", ["a", "b"], 1, tags=["web", "Python"]), taxonomy
    )
    assert _checks(bad) == {"skill_tags"} and not bad[0].critical
    # the same tag gets the same verdict every time it appears
    again = deterministic_findings(
        _row(9, "Other?", ["c", "d"], 2, tags=["web"]), taxonomy
    )
    assert _checks(again) == {"skill_tags"}
    blank = deterministic_findings(_row(2, "Q?", ["a", "b"], 1, tags=[]), taxonomy)
    assert blank[0].extra["tag_status"] == "tags_not_provided"


def test_length_giveaway():
    long = "Last-in first-out ordering used by call stacks and undo history"
    fs = deterministic_findings(
        _row(2, "Which?", [long, "FIFO", "LRU", "Heap"], 1), None
    )
    assert _checks(fs) == {"distractor_parity"} and not fs[0].critical
    assert not deterministic_findings(
        _row(2, "Which?", ["O(1)", "O(n)", "O(log n)", "O(n^2)"], 1), None
    )


def test_image_reference_needs_image_and_normal_text_does_not():
    fs = deterministic_findings(
        _row(2, "Refer to the diagram below. Which node is the root?", ["A", "B"], 1),
        None,
    )
    assert _checks(fs) == {"image_needed"} and fs[0].critical
    assert not deterministic_findings(
        _row(
            2, "Which command builds a Docker image?", ["docker build", "docker run"], 1
        ),
        None,
    )


# ── Stage 1: blind solve vs key ──────────────────────────────────────────────


def _sol(n, correct, total=4, defect="none", detail=""):
    return _Solution(
        row_number=n,
        working="worked it out",
        options=[
            _OptionVerdict(option=i, verdict="correct" if i in correct else "incorrect")
            for i in range(1, total + 1)
        ],
        defect_type=defect,
        defect=detail,
    )


def test_wrong_key_flagged_when_solvers_agree():
    row = _row(2, "2 ** 3 ** 2?", ["64", "512", "256", "8"], 1)
    fs = _solver_findings(row, [_sol(2, {2}), _sol(2, {2})])
    assert fs[0].check == "factual_accuracy" and fs[0].critical
    assert "512" in fs[0].reason and "to 2 (512)" in fs[0].suggestion


def test_solver_disagreement_is_unverified_not_pass_or_fail():
    row = _row(2, "Q?", ["a", "b", "c", "d"], 1)
    fs = _solver_findings(row, [_sol(2, {1}), _sol(2, {3})])
    assert len(fs) == 1 and fs[0].status == "na" and not fs[0].critical
    assert "UNVERIFIED" in fs[0].reason


def test_multiple_correct_and_none_correct():
    row = _row(2, "Which are prime?", ["9", "11", "13", "15"], 2)
    fs = _solver_findings(row, [_sol(2, {2, 3}), _sol(2, {2, 3})])
    assert fs[0].check == "multiple_correct" and fs[0].critical
    none = _solver_findings(row, [_sol(2, set()), _sol(2, set())])
    assert none[0].check == "no_correct_answer" and none[0].critical


def test_correct_key_and_single_solver_ok():
    row = _row(2, "Q?", ["a", "b", "c", "d"], 2)
    assert _solver_findings(row, [_sol(2, {2}), _sol(2, {2})]) == []
    assert _solver_findings(row, [_sol(2, {2}), None]) == []


def test_code_bug_defect():
    row = _row(2, "What does this print?", ["1", "2", "3", "4"], 1)
    bug = _sol(2, {1}, defect="code_bug", detail="off-by-one")
    fs = _solver_findings(row, [bug, bug])
    assert fs[0].check == "code_bug" and fs[0].critical


# ── Stage 3: duplicates ──────────────────────────────────────────────────────


def test_exact_duplicates_ignore_option_order_and_key_number():
    a = _row(2, "Time complexity of binary search?", ["O(n)", "O(log n)", "O(1)"], 2)
    b = _row(3, "time complexity of binary search?", ["O(1)", "O(n)", "O(log n)"], 3)
    ((group, kind),) = _exact_duplicate_groups([a, b])
    assert kind == "duplicate" and {r.row_number for r in group} == {2, 3}


def test_same_stem_same_options_different_answer_is_conflicting_key():
    a = _row(2, "Time complexity of binary search?", ["O(n)", "O(log n)"], 2)
    b = _row(3, "Time complexity of binary search?", ["O(n)", "O(log n)"], 1)
    ((_, kind),) = _exact_duplicate_groups([a, b])
    assert kind == "conflicting_key"


def test_same_stem_different_options_reaches_adjudication_not_auto_flagged():
    a = _row(2, "What is the time complexity of the operation?", ["O(1)", "O(n)"], 1)
    b = _row(
        3, "What is the time complexity of the operation?", ["O(log n)", "O(n^2)"], 1
    )
    assert _exact_duplicate_groups([a, b]) == []
    assert len(_cluster([a, b], None)) == 1


def test_tokenizer_keeps_operators_and_negations():
    toks = svc._token_set("2 ** 3 ** 2 is not true")
    assert {"**", "not", "true"} <= toks


# ── Severity tiers ───────────────────────────────────────────────────────────


def test_minor_only_issues_pass_with_suggestions_and_gtg_otherwise():
    rows = [_row(2, "q1", ["a", "b"], 1), _row(3, "q2", ["a", "b"], 1)]
    findings = {
        2: [
            svc.Finding(
                "grammar_spelling", "fail", "'recieve' is misspelled", "Fix spelling."
            )
        ],
        3: [],
    }
    result = _assemble(rows, findings, True)
    (qr,) = result.question_reviews
    assert qr.verdict == "pass_with_suggestions"
    assert result.row_comments[3] == "gtg" and "recieve" in result.row_comments[2]
    overall = result.set_summary.overall_results
    assert (overall.failed, overall.passed, overall.pass_with_suggestions) == (0, 2, 1)


def test_critical_check_fails_the_row():
    rows = [_row(2, "q1", ["a", "b"], 1)]
    findings = {2: [svc.Finding("factual_accuracy", "fail", "wrong key", "Fix key.")]}
    result = _assemble(rows, findings, True)
    assert result.question_reviews[0].verdict == "fail"
    assert result.set_summary.overall_results.failed == 1


# ── Reviewed workbook ────────────────────────────────────────────────────────


def test_reviewed_workbook_keeps_original_cells_verbatim():
    header = list(EXPECTED_COLUMNS)
    row = ["Q?", "-5", "+3", "=1+1", "0", 1, 1, 0, "Easy", 0, "Python", 0, 0]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(header)
    ws.append(row)
    ws["D2"].data_type = "s"  # a text cell that merely looks like a formula
    buf = io.BytesIO()
    wb.save(buf)
    out = build_reviewed_workbook(buf.getvalue(), "s.xlsx", MCQReviewResult())
    got = openpyxl.load_workbook(io.BytesIO(out)).active
    assert [got.cell(row=2, column=c).value for c in (2, 3, 4)] == ["-5", "+3", "=1+1"]
    assert got.cell(row=2, column=4).data_type == "s"  # stays text, not a formula


# ── End to end with fake providers ───────────────────────────────────────────


async def test_pipeline_end_to_end_with_fake_llms(monkeypatch):
    rows = [
        _row(2, "What is 2 ** 3 ** 2 in Python?", ["64", "512", "256", "8"], 1),
        _row(3, "Which is the cappital of France?", ["Paris", "Rome", "Madrid"], 1),
        _row(4, "What does len([1, 2]) return?", ["1", "2", "3"], 2),
        _row(5, "Name a FIFO structure.", ["Queue", "Banana", "Cloud"], 1),
        _row(6, "Name a FIFO structure.", ["Cloud", "Queue", "Banana"], 2),
    ]
    truth = {2: {2}, 3: {1}, 4: {2}, 5: {1}, 6: {2}}

    async def fake(system: str, user: str) -> str:
        data = json.loads(user)
        if "expert solver" in system:
            return json.dumps(
                {
                    "solutions": [
                        {
                            "row_number": d["row_number"],
                            "working": "w",
                            "options": [
                                {
                                    "option": i,
                                    "verdict": "correct"
                                    if i in truth[d["row_number"]]
                                    else "incorrect",
                                }
                                for i in range(1, len(d["options"]) + 1)
                            ],
                        }
                        for d in data
                    ]
                }
            )
        if "quality auditor" in system:
            reviews = []
            for d in data:
                checks = {}
                if d["row_number"] == 3:
                    checks = {
                        "grammar_spelling": {
                            "status": "fail",
                            "reason": "'cappital' is misspelled",
                        }
                    }
                reviews.append(
                    {
                        "row_number": d["row_number"],
                        "comment": "x",
                        "checks": checks,
                        "suggestion": "Fix spelling." if checks else "",
                    }
                )
            return json.dumps({"question_reviews": reviews})
        raise AssertionError("exact duplicates must not need adjudication")

    async def no_embeddings(_rows):
        return None

    monkeypatch.setitem(svc._PROVIDERS, "openai", fake)
    monkeypatch.setitem(svc._PROVIDERS, "anthropic", fake)
    monkeypatch.setattr(svc, "_embed", no_embeddings)

    result = await review_mcqs(rows, "", {"python"})
    by_row = {qr.row_number: qr for qr in result.question_reviews}

    assert by_row[2].verdict == "fail" and "factual_accuracy" in by_row[2].checks
    assert by_row[3].verdict == "pass_with_suggestions"  # spelling alone doesn't fail
    assert 4 not in by_row and result.row_comments[4] == "gtg"
    assert by_row[5].verdict == by_row[6].verdict == "fail"  # reordered duplicates
    assert by_row[5].checks["duplicate_question"].duplicate_of_rows == [6]
    assert by_row[6].checks["duplicate_question"].duplicate_of_rows == [5]
    assert set(result.row_comments) == {2, 3, 4, 5, 6}  # every row has a comment


# ── Golden set (offline half; the LLM half is scripts/eval_mcq_reviewer.py) ──


def test_golden_set_deterministic_expectations():
    from scripts.eval_mcq_reviewer import TAXONOMY, golden_rows

    rows, expected = golden_rows()
    flagged = {r.row_number: _checks(deterministic_findings(r, TAXONOMY)) for r in rows}
    for n, expect in expected.items():
        stage0 = expect & {"structural_validation", "image_needed", "distractor_parity"}
        if stage0:
            assert stage0 <= flagged[n], f"row {n} should be caught in code"
        if not expect:
            assert not flagged[n], f"good row {n} must not be flagged in code"

    dup_rows = [r for r in rows if r.is_structurally_valid]
    groups = _exact_duplicate_groups(dup_rows)
    dup_expected = {n for n, e in expected.items() if "duplicate_question" in e}
    assert {r.row_number for g, _ in groups for r in g} == dup_expected
